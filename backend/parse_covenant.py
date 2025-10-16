from openpyxl import load_workbook
import re

def _normalize_header(value):
    if value is None:
        return None
    text = str(value).strip().lower()
    # collapse non-alphanumerics to single spaces
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return text.strip()


_HEADER_ALIASES = {
    "calculation date": {"calc date", "calculation date", "date"},
    "show in report": {"show in report", "show", "show in rpt"},
    "sno": {"sno", "s no", "serial", "serial no", "sr no"},
    "covenant name": {"covenant name", "covenant", "name"},
    "threshold": {"threshold"},
    "consequence": {"consequence"},
    "borrower calculation": {"borrower calculation", "borrower", "borrower calc"},
    "lender calculation": {"lender calculation", "lender", "lender calc"},
    "compliance check": {"compliance check", "compliance", "check"},
    "comment": {"comment", "remarks", "notes"},
    "source file": {"source file", "source", "src"},
    "reference file": {"reference file", "reference", "ref"},
}


def _match_headers(row_values):
    headers_normalized = [_normalize_header(c) for c in row_values]
    header_map = {}
    for idx, norm in enumerate(headers_normalized):
        if not norm:
            continue
        for canonical, variants in _HEADER_ALIASES.items():
            if norm == canonical or norm in variants:
                # prefer first occurrence
                header_map.setdefault(canonical, idx)
    return header_map


def parse_xlsx(file_path):
    wb = load_workbook(file_path)
    sheet = wb.active

    data = []
    current_entry = {}

    # Identify header row automatically (by presence of a recognizable covenant header)
    header_row_idx = None
    header_map = {}
    for i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if not row:
            continue
        tentative_map = _match_headers(row)
        if "covenant name" in tentative_map:
            header_row_idx = i
            header_map = tentative_map
            break
    if header_row_idx is None:
        raise ValueError("No header row found (missing a 'Covenant Name' column)")

    for row in sheet.iter_rows(min_row=header_row_idx+1, values_only=True):
        # Normalize row
        row = list(row) if row else []

        def get_value(canonical, default_index=None):
            idx = header_map.get(canonical)
            if idx is None:
                if default_index is None:
                    return None
                idx = default_index
            return row[idx] if idx is not None and idx < len(row) else None

        calc_date = get_value("calculation date", 0)
        show_in_report = get_value("show in report", 1)
        sno = get_value("sno", 2)
        covenant_name = get_value("covenant name", 3)
        threshold = get_value("threshold", 4)
        consequence = get_value("consequence", 5)
        borrower_calc = get_value("borrower calculation", 6)
        lender_calc = get_value("lender calculation", 7)
        compliance_check = get_value("compliance check", 8)
        comment = get_value("comment", 9)
        source_file = get_value("source file", 10)
        reference_file = get_value("reference file", 11)

        # Start a new block if SNO present
        if sno is not None:
            if current_entry:
                data.append(current_entry)
            current_entry = {
                "Calculation Date": calc_date,
                "Show in Report": show_in_report,
                "SNO": sno,
                "Covenants": []
            }

        # Append covenant details
        if covenant_name:
            if not current_entry:
                current_entry = {
                    "Calculation Date": calc_date,
                    "Show in Report": show_in_report,
                    "SNO": None,
                    "Covenants": []
                }
            current_entry["Covenants"].append({
                "Covenant Name": covenant_name,
                "Threshold": threshold,
                "Consequence": consequence,
                "Borrower Calculation": borrower_calc,
                "Lender Calculation": lender_calc,
                "Compliance Check": compliance_check,
                "Comment": comment,
                "Source File": source_file,
                "Reference File": reference_file
            })

    if current_entry:
        data.append(current_entry)

    return data


# Minimal placeholder to satisfy imports (PDF parsing not implemented here)
def parse_compliance_certificate(file_path):
    return {"calc_date": None, "rows": [], "items": []}